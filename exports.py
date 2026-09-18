import csv
import io
import re
import time
from urllib.parse import quote_plus

DISPLAY_COLUMNS = [
    ("name", "Name"),
    ("address", "Adresse"),
    ("phone", "Telefon"),
    ("website", "Website"),
    ("google_search_url", "Google Profil (Suche)"),
    ("icp_score", "ICP-Score"),
    ("icp_tier", "ICP-Tier"),
    ("competitor_system", "Erkanntes System"),
    ("score", "Konkurrenz-Ampel"),
    ("rating_count", "Anzahl Bewertungen"),
    ("open_now", "Geoeffnet"),
    ("opening_status", "Eroeffnungsstatus"),
    ("opening_date", "Geplantes Eroeffnungsdatum"),
    ("chain_flag", "Filialkette (3-9)"),
    ("pain_points", "Review Pain-Points"),
    ("status", "Status"),
    ("notes", "Notiz"),
    ("region_query", "Gesuchte Region"),
    ("category_query", "Gesuchte Kategorie"),
    ("first_seen_display", "Erstmals gefunden"),
]

FIRST_SEEN_RE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})(?: (\d{2}):(\d{2}))?")


def google_search_url(name: str, address: str) -> str:
    query = f"{name or ''} {address or ''}".strip()
    return f"https://www.google.com/search?q={quote_plus(query)}" if query else ""


def format_first_seen(raw: str) -> str:
    """Wandelt den intern sortierbaren Zeitstempel in 'HH:MM - TT-MM-JJJJ' um (CET/CEST)."""
    if not raw:
        return ""
    m = FIRST_SEEN_RE.match(raw)
    if not m:
        return raw
    y, mo, d, h, mi = m.groups()
    return f"{h}:{mi} - {d}-{mo}-{y}" if h else f"{d}-{mo}-{y}"


def _with_computed_fields(rows: list) -> list:
    return [{**r, "google_search_url": google_search_url(r.get("name"), r.get("address")),
              "first_seen_display": format_first_seen(r.get("first_seen"))} for r in rows]


def rows_to_csv_bytes(rows: list) -> bytes:
    rows = _with_computed_fields(rows)
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow([label for _, label in DISPLAY_COLUMNS])
    for r in rows:
        writer.writerow([r.get(key, "") for key, _ in DISPLAY_COLUMNS])
    return buf.getvalue().encode("utf-8-sig")


def rows_to_xlsx_bytes(rows: list) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font

    rows = _with_computed_fields(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append([label for _, label in DISPLAY_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    fills = {
        "A": PatternFill(start_color="FF9C9C", end_color="FF9C9C", fill_type="solid"),
        "B": PatternFill(start_color="FFE59C", end_color="FFE59C", fill_type="solid"),
        "C": PatternFill(start_color="D6F0D6", end_color="D6F0D6", fill_type="solid"),
    }
    pain_fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
    tier_col_idx = [k for k, _ in DISPLAY_COLUMNS].index("icp_tier") + 1
    pain_col_idx = [k for k, _ in DISPLAY_COLUMNS].index("pain_points") + 1

    for row_i, r in enumerate(rows, start=2):
        for col_i, (key, _) in enumerate(DISPLAY_COLUMNS, start=1):
            ws.cell(row=row_i, column=col_i, value=r.get(key, ""))
        tier = r.get("icp_tier", "")
        if tier in fills:
            ws.cell(row=row_i, column=tier_col_idx).fill = fills[tier]
        if r.get("pain_points"):
            ws.cell(row=row_i, column=pain_col_idx).fill = pain_fill

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 45)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rows_to_hubspot_csv_bytes(rows: list) -> bytes:
    rows = _with_computed_fields(rows)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Company name", "Phone Number", "Website URL", "Google Profil (Suche)", "Street Address",
                      "ICP Score", "ICP Tier", "Erkanntes Konkurrenzsystem", "Status", "Notiz"])
    for r in rows:
        note_parts = []
        if r.get("opening_status") and r.get("opening_status") != "Etabliert":
            date_part = f" ({r['opening_date']})" if r.get("opening_date") else ""
            note_parts.append(f"{r['opening_status']}{date_part}")
        if r.get("chain_flag"):
            note_parts.append("Teil einer Filialkette (3-9 Standorte)")
        if r.get("pain_points"):
            note_parts.append(f"Review Pain-Points: {r['pain_points']}")
        if r.get("notes"):
            note_parts.append(r["notes"])
        writer.writerow([
            r.get("name", ""), r.get("phone", ""), r.get("website", ""), r.get("google_search_url", ""),
            r.get("address", ""), r.get("icp_score", ""), r.get("icp_tier", ""), r.get("competitor_system", ""),
            r.get("status", ""), " | ".join(note_parts),
        ])
    return buf.getvalue().encode("utf-8-sig")


def timestamped(name: str, ext: str) -> str:
    return f"{name}_{time.strftime('%Y%m%d_%H%M%S')}.{ext}"
